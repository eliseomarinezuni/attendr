from __future__ import annotations

import sys
import tempfile
import unittest
from io import BytesIO
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from canvasapi.exceptions import CanvasException, InvalidAccessToken
from pptx import Presentation

from academic_assistant.canvas_client import (
    CanvasAuthenticationError,
    CanvasClient,
    CanvasConfigurationError,
)

UTC = timezone.utc
NOW = datetime(2026, 9, 4, 12, 0, tzinfo=UTC)


class FakeCourse(SimpleNamespace):
    def get_assignments(self, **kwargs):
        if getattr(self, "assignment_error", None):
            raise self.assignment_error
        self.assignment_kwargs = kwargs
        return iter(getattr(self, "assignments", []))

    def get_discussion_topics(self, **kwargs):
        if getattr(self, "announcement_error", None):
            raise self.announcement_error
        self.announcement_kwargs = kwargs
        return iter(getattr(self, "announcements", []))

    def get_files(self, **kwargs):
        self.file_kwargs = kwargs
        return iter(getattr(self, "files", []))

    def get_modules(self, **kwargs):
        if getattr(self, "module_error", None):
            raise self.module_error
        self.module_kwargs = kwargs
        return iter(getattr(self, "modules", []))

    def get_file(self, file_id, **kwargs):
        direct = getattr(self, "direct_files", {})
        if str(file_id) in direct:
            return direct[str(file_id)]
        return next(item for item in self.files if str(item.id) == str(file_id))

    def get_page(self, page_url, **kwargs):
        return self.pages[page_url]

    def get_pages(self, **kwargs):
        if getattr(self, "page_error", None):
            raise self.page_error
        self.pages_kwargs = kwargs
        return iter(getattr(self, "page_summaries", []))


class FakeFile(SimpleNamespace):
    def get_contents(self, *, binary=False):
        self.binary_requested = binary
        return self.content


class FakeCanvas:
    def __init__(self, courses, events=None, event_errors=None, user_error=None):
        self.courses = courses
        self.events = events or {}
        self.event_errors = event_errors or {}
        self.user_error = user_error
        self.course_kwargs = None

    def get_current_user(self):
        if self.user_error:
            raise self.user_error
        return SimpleNamespace(id=42, name="Ada Student")

    def get_courses(self, **kwargs):
        self.course_kwargs = kwargs
        return iter(self.courses)

    def get_calendar_events(self, **kwargs):
        course_id = int(kwargs["context_codes"][0].split("_")[1])
        if course_id in self.event_errors:
            raise self.event_errors[course_id]
        return iter(self.events.get(course_id, []))


def assignment(item_id, title, due_at, **overrides):
    values = {
        "id": item_id,
        "name": title,
        "due_at": due_at,
        "html_url": f"https://canvas.example/courses/1/assignments/{item_id}",
        "updated_at": "2026-09-01T10:00:00Z",
        "points_possible": 10,
        "submission_types": ["online_upload"],
        "description": "<p>Instructions</p>",
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def event(item_id, title, start_at, **overrides):
    values = {
        "id": item_id,
        "title": title,
        "start_at": start_at,
        "end_at": None,
        "all_day": False,
        "html_url": f"https://canvas.example/calendar?event_id={item_id}",
        "updated_at": "2026-09-02T10:00:00Z",
        "description": None,
    }
    values.update(overrides)
    return SimpleNamespace(**values)


def announcement(item_id, title, posted_at, read_state="unread"):
    return SimpleNamespace(
        id=item_id,
        title=title,
        posted_at=posted_at,
        read_state=read_state,
        message="<p>Hello <strong>class</strong>.</p><script>ignore me</script>",
        html_url=f"https://canvas.example/announcements/{item_id}",
        author={"display_name": "Professor Example"},
    )


def course(course_id=1, name="Algorithms", **overrides):
    values = {
        "id": course_id,
        "name": name,
        "course_code": "CSC300",
        "html_url": f"https://canvas.example/courses/{course_id}",
        "assignments": [],
        "announcements": [],
        "files": [],
        "modules": [],
        "pages": {},
        "page_summaries": [],
        "direct_files": {},
        "syllabus_body": "",
    }
    values.update(overrides)
    return FakeCourse(**values)


class CanvasClientTests(unittest.TestCase):
    def make_client(self, fake_canvas, **overrides):
        values = {
            "lookahead_days": 30,
            "announcement_days": 14,
            "app_timezone": "America/Toronto",
            "canvas": fake_canvas,
            "now_provider": lambda: NOW,
        }
        values.update(overrides)
        return CanvasClient(
            "https://canvas.example",
            "secret-token",
            **values,
        )

    def test_snapshot_normalizes_items_ids_types_and_timezone(self):
        active_course = course(
            assignments=[
                assignment(10, "Midterm Exam", "2026-09-10T18:00:00Z"),
                assignment(
                    11,
                    "Chapter Quiz",
                    "2026-09-11T18:00:00Z",
                    submission_types=["online_quiz"],
                ),
                assignment(12, "Problem Set", "2026-09-12T18:00:00Z"),
            ],
        )
        fake = FakeCanvas(
            [active_course],
            events={1: [event(20, "Final review session", "2026-09-13T18:00:00Z")]},
        )

        snapshot = self.make_client(fake).fetch_snapshot()

        self.assertEqual(snapshot.user_name, "Ada Student")
        self.assertEqual(
            [item.kind for item in snapshot.items],
            ["exam", "quiz", "assignment", "exam"],
        )
        self.assertEqual(snapshot.items[0].uid, "canvas:assignment:1:10")
        self.assertEqual(snapshot.items[-1].uid, "canvas:event:1:20")
        self.assertEqual(snapshot.items[0].due_at_local.hour, 14)
        self.assertEqual(snapshot.items[0].due_at_local.tzname(), "EDT")
        self.assertEqual(fake.course_kwargs["enrollment_state"], "active")
        self.assertEqual(fake.course_kwargs["enrollment_type"], "student")
        self.assertTrue(active_course.assignment_kwargs["override_assignment_dates"])
        self.assertEqual(active_course.assignment_kwargs["include"], ["submission"])

    def test_downloads_syllabus_page_named_and_linked_pdfs_only(self):
        pdf = b"%PDF-1.4\nminimal test content"
        active_course = course(
            syllabus_body='<p>Exam Oct 20</p><a href="/courses/1/files/12">outline</a>',
            files=[
                FakeFile(
                    id=11,
                    display_name="CSCI_3101U_F26_Syllabus.pdf",
                    content_type="application/pdf",
                    size=len(pdf),
                    content=pdf,
                    updated_at="2026-09-01T10:00:00Z",
                ),
                FakeFile(
                    id=12,
                    display_name="CSC 3000.pdf",
                    content_type="application/pdf",
                    size=len(pdf),
                    content=pdf,
                    updated_at="2026-09-01T10:00:00Z",
                ),
                FakeFile(
                    id=13,
                    display_name="Lecture 1.pdf",
                    content_type="application/pdf",
                    size=len(pdf),
                    content=pdf,
                    updated_at="2026-09-01T10:00:00Z",
                ),
            ],
        )
        client = self.make_client(FakeCanvas([active_course]))

        with tempfile.TemporaryDirectory() as directory:
            report = client.download_syllabus_materials(directory)

            self.assertEqual(len(report.materials), 3)
            self.assertTrue(all(item.local_path.is_file() for item in report.materials))
            titles = {item.title for item in report.materials}
            self.assertIn("Canvas syllabus page", titles)
            self.assertIn("CSCI_3101U_F26_Syllabus.pdf", titles)
            self.assertIn("CSC 3000.pdf", titles)
            self.assertNotIn("Lecture 1.pdf", titles)
            self.assertEqual(
                active_course.file_kwargs["content_types"],
                [
                    "application/pdf",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ],
            )

    def test_directly_downloads_linked_file_hidden_from_files_area(self):
        pdf = b"%PDF-1.4\nhidden but linked syllabus"
        linked = FakeFile(
            id=44,
            display_name="CourseOutline.pdf",
            content_type="application/pdf",
            size=len(pdf),
            content=pdf,
            updated_at="2026-09-01T10:00:00Z",
            hidden_for_user=True,
            locked=False,
        )
        active_course = course(
            syllabus_body='<a href="/courses/1/files/44">Syllabus</a>',
            direct_files={"44": linked},
        )
        with tempfile.TemporaryDirectory() as directory:
            report = self.make_client(FakeCanvas([active_course])).download_syllabus_materials(directory)
        self.assertEqual(len(report.materials), 2)
        self.assertTrue(any(item.source_id == "44" for item in report.materials))

    def test_finds_generic_pdf_linked_from_syllabus_module(self):
        pdf = b"%PDF-1.4\ncourse dates"
        generic = FakeFile(
            id=31,
            display_name="document.pdf",
            content_type="application/pdf",
            size=len(pdf),
            content=pdf,
            updated_at="2026-09-01T10:00:00Z",
        )
        active_course = course(
            files=[generic],
            modules=[
                SimpleNamespace(
                    name="Course Syllabus",
                    locked_for_user=False,
                    items=[
                        {
                            "type": "File",
                            "content_id": 31,
                            "title": "Download document",
                        }
                    ],
                )
            ],
        )
        with tempfile.TemporaryDirectory() as directory:
            report = self.make_client(FakeCanvas([active_course])).download_syllabus_materials(directory)
        self.assertEqual(len(report.materials), 1)
        self.assertEqual(report.materials[0].title, "document.pdf")

    def test_finds_syllabus_content_inside_generically_named_page(self):
        page = SimpleNamespace(
            url="home",
            title="Home",
            body="<p>Course syllabus: Midterm October 20, 2026</p>",
            updated_at="2026-09-01T10:00:00Z",
        )
        active_course = course(
            pages={"home": page},
            page_summaries=[SimpleNamespace(url="home", title="Home")],
        )
        with tempfile.TemporaryDirectory() as directory:
            report = self.make_client(FakeCanvas([active_course])).download_syllabus_materials(directory)
        self.assertEqual(len(report.materials), 1)
        self.assertEqual(report.materials[0].content_type, "text/html")

    def test_missing_syllabus_areas_are_nonfatal(self):
        active_course = course(
            module_error=CanvasException("blocked"),
            page_error=CanvasException("blocked"),
        )
        with tempfile.TemporaryDirectory() as directory:
            report = self.make_client(FakeCanvas([active_course])).download_syllabus_materials(directory)
        self.assertEqual(report.materials, ())
        self.assertEqual(len(report.warnings), 2)

    def test_downloads_lecture_module_material_but_excludes_lab(self):
        pdf = b"%PDF-1.4\nlecture content"
        lecture_file = FakeFile(
            id=21,
            display_name="Lecture 1.pdf",
            content_type="application/pdf",
            size=len(pdf),
            content=pdf,
            updated_at="2026-09-08T10:00:00Z",
        )
        lab_file = FakeFile(
            id=22,
            display_name="Lab 1.pdf",
            content_type="application/pdf",
            size=len(pdf),
            content=pdf,
            updated_at="2026-09-08T10:00:00Z",
        )
        module = SimpleNamespace(
            name="Week 1",
            position=1,
            locked_for_user=False,
            items=[
                {"type": "File", "content_id": 21, "title": "Lecture 1", "position": 1},
                {"type": "File", "content_id": 22, "title": "Lab 1", "position": 2},
            ],
        )
        active_course = course(files=[lecture_file, lab_file], modules=[module])
        with tempfile.TemporaryDirectory() as directory:
            report = self.make_client(FakeCanvas([active_course])).download_lecture_materials(directory)
        self.assertEqual([item.title for item in report.materials], ["Lecture 1"])
        self.assertEqual(report.materials[0].module_name, "Week 1")

    def test_downloads_first_lecture_from_introduction_module(self):
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[1])
        output = BytesIO()
        presentation.save(output)
        content = output.getvalue()
        introduction = FakeFile(
            id=23,
            display_name="Introduction_canvas.pptx",
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            size=len(content),
            content=content,
            updated_at="2026-09-09T10:00:00Z",
        )
        module = SimpleNamespace(
            name="Introduction",
            position=4,
            locked_for_user=False,
            items=[
                {
                    "type": "File",
                    "content_id": 23,
                    "title": "Introduction_canvas.pptx",
                    "position": 4,
                }
            ],
        )
        active_course = course(files=[introduction], modules=[module])
        with tempfile.TemporaryDirectory() as directory:
            report = self.make_client(FakeCanvas([active_course])).download_lecture_materials(directory)
        self.assertEqual([item.title for item in report.materials], ["Introduction_canvas.pptx"])
        self.assertEqual(report.materials[0].module_name, "Introduction")

    def test_upcoming_window_includes_boundaries_and_excludes_invalid_dates(self):
        active_course = course(
            assignments=[
                assignment(1, "Starts now", "2026-09-04T12:00:00Z"),
                assignment(2, "Ends window", "2026-10-04T12:00:00Z"),
                assignment(3, "Too late", "2026-10-04T12:00:01Z"),
                assignment(4, "Undated", None),
            ]
        )
        snapshot = self.make_client(FakeCanvas([active_course])).fetch_snapshot()

        self.assertEqual([item.source_id for item in snapshot.items], ["1", "2"])
        self.assertTrue(
            any(
                "Skipped undated assignment" in warning
                for warning in snapshot.warnings
            )
        )

    def test_all_day_event_uses_local_midnight(self):
        active_course = course()
        fake = FakeCanvas(
            [active_course],
            events={
                1: [
                    event(
                        5,
                        "Department event",
                        None,
                        all_day=True,
                        all_day_date="2026-09-20",
                    )
                ]
            },
        )

        item = self.make_client(fake).fetch_snapshot().items[0]

        self.assertTrue(item.all_day)
        self.assertEqual(item.due_at_local.isoformat(), "2026-09-20T00:00:00-04:00")

    def test_announcements_are_recent_unread_and_plain_text(self):
        active_course = course(
            announcements=[
                announcement(1, "Recent", "2026-09-03T12:00:00Z"),
                announcement(2, "Old", "2026-08-20T11:59:59Z"),
                announcement(3, "Already read", "2026-09-02T12:00:00Z", "read"),
            ]
        )

        snapshot = self.make_client(FakeCanvas([active_course])).fetch_snapshot()

        self.assertEqual([item.title for item in snapshot.announcements], ["Recent"])
        self.assertEqual(snapshot.announcements[0].message_text, "Hello class.")
        self.assertEqual(
            active_course.announcement_kwargs,
            {
                "only_announcements": True,
                "filter_by": "unread",
                "order_by": "recent_activity",
            },
        )

    def test_announcement_uses_created_timestamp_when_posted_timestamp_is_null(self):
        item = announcement(1, "Delayed announcement", None)
        item.created_at = "2026-09-03T12:00:00Z"
        snapshot = self.make_client(
            FakeCanvas([course(announcements=[item])])
        ).fetch_snapshot()

        self.assertTrue(snapshot.complete)
        self.assertEqual(
            snapshot.announcements[0].posted_at.isoformat(),
            "2026-09-03T12:00:00+00:00",
        )

    def test_one_inaccessible_course_adds_warnings_without_losing_other_data(self):
        good = course(
            1,
            "Accessible",
            assignments=[assignment(1, "Homework", "2026-09-05T12:00:00Z")],
            announcements=[announcement(1, "News", "2026-09-03T12:00:00Z")],
        )
        bad = course(
            2,
            "Restricted",
            assignment_error=CanvasException("blocked"),
            announcement_error=CanvasException("blocked"),
        )
        fake = FakeCanvas(
            [good, bad],
            event_errors={2: CanvasException("blocked")},
        )

        snapshot = self.make_client(fake).fetch_snapshot()

        self.assertEqual(len(snapshot.items), 1)
        self.assertEqual(len(snapshot.announcements), 1)
        self.assertEqual(len(snapshot.warnings), 3)
        self.assertTrue(all("Restricted" in warning for warning in snapshot.warnings))
        self.assertEqual(snapshot.incomplete_course_ids, (2,))

    def test_invalid_token_becomes_actionable_authentication_error(self):
        fake = FakeCanvas([], user_error=InvalidAccessToken("bad token"))
        with self.assertRaisesRegex(
            CanvasAuthenticationError, "rejected the API token"
        ):
            self.make_client(fake).validate_credentials()

    def test_invalid_configuration_fails_before_any_request(self):
        with self.assertRaisesRegex(CanvasConfigurationError, "origin only"):
            CanvasClient("https://canvas.example/api/v1", "token")
        with self.assertRaisesRegex(CanvasConfigurationError, "CANVAS_API_TOKEN"):
            CanvasClient("https://canvas.example", "")
        with self.assertRaisesRegex(CanvasConfigurationError, "IANA timezone"):
            CanvasClient("https://canvas.example", "token", app_timezone="Mars/Olympus")


if __name__ == "__main__":
    unittest.main()
