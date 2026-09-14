from __future__ import annotations

import sys
import json
import tempfile
import unittest
from dataclasses import replace
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from unittest.mock import patch

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from academic_assistant import (
    AIProviderError,
    CourseSchedule,
    CourseMaterialsSync,
    MaterialDownloadReport,
    SyllabusMaterial,
)

UTC = timezone.utc
NOW = datetime(2026, 9, 5, 12, 0, tzinfo=UTC)


class FakeCanvas:
    def __init__(self, materials, warnings=(), incomplete_course_ids=()):
        self.materials = materials
        self.warnings = warnings
        self.incomplete_course_ids = incomplete_course_ids

    def download_syllabus_materials(
        self, directory, *, max_file_bytes, known_file_ids_by_course=None, course_ids=None
    ):
        return MaterialDownloadReport(
            tuple(self.materials), tuple(self.warnings), tuple(self.incomplete_course_ids)
        )


class FakeAI:
    def __init__(self, deadlines):
        self.deadlines = deadlines
        self.calls = []

    def extract_major_deadlines(self, text, **kwargs):
        self.calls.append((text, kwargs))
        return self.deadlines


def deadline(
    date_value="2026-10-20",
    time_value="13:30",
    *,
    title="Midterm Exam",
    evidence="Midterm Exam: October 20 at 1:30 PM",
):
    return {
        "title": title,
        "due_date": date_value,
        "due_time": time_value,
        "kind": "exam",
        "source_evidence": evidence,
    }


class CourseMaterialsSyncTests(unittest.TestCase):
    def make_material(self, path, *, uid="canvas:syllabus-page:1", digest=None):
        content = path.read_bytes()
        return SyllabusMaterial(
            uid=uid,
            source_id="syllabus-page",
            course_id=1,
            course_name="Algorithms",
            title="Canvas syllabus page",
            content_type="text/html",
            local_path=path,
            content_sha256=digest or sha256(content).hexdigest(),
            updated_at=None,
            html_url="https://canvas.example/courses/1/assignments/syllabus",
        )

    def test_changed_material_is_analyzed_then_unchanged_hash_uses_cache(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            html_path = root / "syllabus.html"
            html_path.write_text("<p>Midterm Exam: October 20, 2026 at 1:30 PM</p>")
            material = self.make_material(html_path)
            ai = FakeAI([deadline()])
            first = CourseMaterialsSync(
                FakeCanvas([material]),
                ai,
                materials_directory=root / "materials",
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync()

            cached_ai = FakeAI([])
            second = CourseMaterialsSync(
                FakeCanvas([material]),
                cached_ai,
                materials_directory=root / "materials",
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync()

            self.assertEqual(first.materials_analyzed, 1)
            self.assertEqual(len(first.items), 1)
            self.assertEqual(first.items[0].kind, "exam")
            self.assertEqual(
                first.items[0].due_at_local.isoformat(), "2026-10-20T13:30:00-04:00"
            )
            self.assertEqual(second.cached_materials_reused, 1)
            self.assertEqual(cached_ai.calls, [])
            self.assertEqual(second.items[0].uid, first.items[0].uid)

    def test_schema_valid_hallucinated_date_never_becomes_academic_item(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            html_path = root / "syllabus.html"
            html_path.write_text("<p>Midterm Exam date will be announced later.</p>")
            material = self.make_material(html_path)
            hallucination = {
                "title": "Midterm Exam",
                "due_date": "2026-10-20",
                "due_time": None,
                "kind": "exam",
                "source_evidence": "Midterm Exam date will be announced later.",
            }

            report = CourseMaterialsSync(
                FakeCanvas([material]),
                FakeAI([hallucination]),
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync()

        self.assertEqual(report.items, ())
        self.assertTrue(report.complete)
        self.assertEqual(report.blocking_warnings, ())
        self.assertTrue(any("Rejected ungrounded AI deadline" in warning for warning in report.warnings))

    def test_legacy_unverified_cache_is_revalidated_on_provider_failure(self):
        class FailingAI:
            def extract_major_deadlines(self, *args, **kwargs):
                raise AIProviderError("provider unavailable")

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            html_path = root / "syllabus.html"
            html_path.write_text("<p>Midterm Exam: October 20, 2026</p>")
            material = self.make_material(html_path)
            index_path = root / "index.json"
            index_path.write_text(json.dumps({
                "version": 1,
                "sources": {
                    material.uid: {
                        "content_sha256": material.content_sha256,
                        "course_id": material.course_id,
                        "course_name": material.course_name,
                        "title": material.title,
                        "deadlines": [deadline(
                            "2026-10-21", None,
                            evidence="Midterm Exam: October 20, 2026",
                        )],
                        "extraction_version": 3,
                        "context_hash": "legacy",
                    }
                },
            }))

            report = CourseMaterialsSync(
                FakeCanvas([material]),
                FailingAI(),
                index_path=index_path,
                now_provider=lambda: NOW,
            ).sync()

        self.assertEqual(report.items, ())
        self.assertFalse(report.complete)
        self.assertTrue(report.blocking_warnings)
        self.assertTrue(any("Rejected ungrounded AI deadline" in warning for warning in report.warnings))

    def test_unchanged_legacy_cache_is_grounded_and_upgraded_without_gemini(self):
        class ForbiddenAI:
            model = "test-model"

            def extract_major_deadlines(self, *args, **kwargs):
                raise AssertionError("Gemini must not be called")

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            html_path = root / "syllabus.html"
            html_path.write_text("<p>Midterm Exam: October 20, 2026</p>")
            material = self.make_material(html_path)
            index_path = root / "index.json"
            index_path.write_text(json.dumps({
                "version": 1,
                "sources": {
                    material.uid: {
                        "content_sha256": material.content_sha256,
                        "course_id": material.course_id,
                        "course_name": material.course_name,
                        "title": material.title,
                        "deadlines": [deadline("2026-10-20", None)],
                        "extraction_version": 3,
                        "context_hash": "legacy",
                    }
                },
            }))

            report = CourseMaterialsSync(
                FakeCanvas([material]),
                ForbiddenAI(),
                index_path=index_path,
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})
            upgraded = json.loads(index_path.read_text())["sources"][material.uid]

        self.assertTrue(report.complete)
        self.assertEqual(report.cached_materials_reused, 1)
        self.assertEqual(len(report.items), 1)
        self.assertEqual(upgraded["extraction_version"], 5)
        self.assertNotEqual(upgraded["context_hash"], "legacy")

    def test_changed_source_provider_outage_is_blocking_and_preserves_old_cache(self):
        class FailingAI:
            def extract_major_deadlines(self, *args, **kwargs):
                raise AIProviderError("provider unavailable", transient=True)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path = root / "old.html"
            old_path.write_text("<p>Midterm Exam: October 20, 2026</p>")
            old = self.make_material(old_path)
            index_path = root / "index.json"
            CourseMaterialsSync(
                FakeCanvas([old]),
                FakeAI([deadline("2026-10-20", None)]),
                index_path=index_path,
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})
            original_index = index_path.read_text()

            changed_path = root / "changed.html"
            changed_path.write_text("<p>Midterm Exam: October 27, 2026</p>")
            changed = self.make_material(changed_path)
            report = CourseMaterialsSync(
                FakeCanvas([changed]),
                FailingAI(),
                index_path=index_path,
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})
            preserved_index = index_path.read_text()

        self.assertFalse(report.complete)
        self.assertTrue(report.blocking_warnings)
        self.assertEqual(report.items, ())
        self.assertEqual(preserved_index, original_index)

    def test_identical_content_at_two_canvas_paths_uses_one_extraction(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first_path = root / "one.html"
            second_path = root / "two.html"
            body = "<p>Midterm Exam: October 20, 2026 at 1:30 PM</p>"
            first_path.write_text(body)
            second_path.write_text(body)
            first = self.make_material(first_path, uid="canvas:syllabus-page:1")
            second = self.make_material(second_path, uid="canvas:syllabus-file:1:9")
            ai = FakeAI([deadline()])

            report = CourseMaterialsSync(
                FakeCanvas([first, second]),
                ai,
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})

        self.assertEqual(len(ai.calls), 1)
        self.assertEqual(report.materials_analyzed, 1)
        self.assertEqual(report.cached_materials_reused, 1)
        self.assertEqual(len(report.items), 1)

    def test_grounding_version_change_reverifies_without_gemini(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "syllabus.html"
            path.write_text("<p>Midterm Exam: October 20, 2026 at 1:30 PM</p>")
            material = self.make_material(path)
            index_path = root / "index.json"
            CourseMaterialsSync(
                FakeCanvas([material]),
                FakeAI([deadline()]),
                index_path=index_path,
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})
            forbidden = FakeAI([])
            with patch("academic_assistant.materials_sync.GROUNDING_VERSION", 3):
                report = CourseMaterialsSync(
                    FakeCanvas([material]),
                    forbidden,
                    index_path=index_path,
                    now_provider=lambda: NOW,
                ).sync(active_course_ids={1})

        self.assertEqual(forbidden.calls, [])
        self.assertTrue(report.complete)

    def test_prompt_version_change_regenerates_only_affected_extraction(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "syllabus.html"
            path.write_text("<p>Midterm Exam: October 20, 2026 at 1:30 PM</p>")
            material = self.make_material(path)
            index_path = root / "index.json"
            CourseMaterialsSync(
                FakeCanvas([material]),
                FakeAI([deadline()]),
                index_path=index_path,
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})
            regenerated = FakeAI([deadline()])
            with patch("academic_assistant.materials_sync.AI_EXTRACTION_VERSION", 6):
                report = CourseMaterialsSync(
                    FakeCanvas([material]),
                    regenerated,
                    index_path=index_path,
                    now_provider=lambda: NOW,
                ).sync(active_course_ids={1})

        self.assertEqual(len(regenerated.calls), 1)
        self.assertTrue(report.complete)

    def test_untimed_midterm_uses_matching_lecture_slot(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            html_path = root / "syllabus.html"
            html_path.write_text("<p>Midterm Exam October 27, 2026</p>")
            material = replace(
                self.make_material(html_path),
                course_name="202699 - Web Dev - EXMP-3030",
            )
            report = CourseMaterialsSync(
                FakeCanvas([material]),
                FakeAI(
                    [
                        deadline(
                            "2026-10-27",
                            None,
                            evidence="Midterm Exam October 27, 2026",
                        )
                    ]
                ),
                index_path=root / "index.json",
                course_schedule=CourseSchedule.load(
                    PROJECT_ROOT / "data/course_schedule.json"
                ),
                now_provider=lambda: NOW,
            ).sync()

            self.assertEqual(
                report.items[0].due_at_local.isoformat(),
                "2026-10-27T12:40:00-04:00",
            )
            self.assertEqual(
                report.items[0].end_at.astimezone(report.items[0].due_at_local.tzinfo).strftime("%H:%M"),
                "14:00",
            )

    def test_conflicting_dates_for_same_exam_are_not_sent_to_calendar(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first_path = root / "one.html"
            second_path = root / "two.html"
            first_path.write_text("<p>Midterm Exam October 20</p>")
            second_path.write_text("<p>Midterm Exam October 21</p>")
            first = self.make_material(first_path, uid="canvas:syllabus-page:1")
            second = self.make_material(second_path, uid="canvas:syllabus-file:1:2")

            class ConflictingAI:
                def extract_major_deadlines(self, text, **kwargs):
                    day = "21" if "21" in text else "20"
                    return [
                        deadline(
                            f"2026-10-{day}",
                            None,
                            evidence=f"Midterm Exam October {day}",
                        )
                    ]

            report = CourseMaterialsSync(
                FakeCanvas([first, second]),
                ConflictingAI(),
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync()

            self.assertEqual(report.items, ())
            self.assertTrue(
                any("Conflicting syllabus dates" in item for item in report.warnings)
            )
            self.assertFalse(report.complete)
            self.assertTrue(report.blocking_warnings)

    def test_word_syllabus_extracts_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "outline.docx"
            document = Document()
            document.add_paragraph("Course schedule")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Assignment 1"
            table.cell(0, 1).text = "October 16, 2026"
            document.save(path)
            material = SyllabusMaterial(
                uid="canvas:syllabus-file:1:9",
                source_id="9",
                course_id=1,
                course_name="Algorithms",
                title="outline.docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                local_path=path,
                content_sha256=sha256(path.read_bytes()).hexdigest(),
                updated_at=None,
                html_url=None,
            )
            ai = FakeAI([deadline(
                "2026-10-16",
                None,
                title="Assignment 1",
                evidence="Assignment 1 | October 16, 2026",
            )])
            report = CourseMaterialsSync(
                FakeCanvas([material]),
                ai,
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync()
        self.assertIn("Assignment 1 | October 16, 2026", ai.calls[0][0])
        self.assertEqual(report.items[0].due_at_local.isoformat(), "2026-10-16T00:00:00-04:00")

    def test_optional_route_failure_is_complete_when_course_has_material(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "syllabus.html"
            path.write_text("<p>Course outline</p>")
            material = self.make_material(path)
            report = CourseMaterialsSync(
                FakeCanvas(
                    [material],
                    warnings=("Could not retrieve syllabus files",),
                    incomplete_course_ids=(1,),
                ),
                FakeAI([]),
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})

        self.assertTrue(report.complete)

    def test_hidden_files_tab_retains_missing_cached_source_without_blocking(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path = root / "old.html"
            old_path.write_text("<p>Midterm Exam: October 20, 2026 at 1:30 PM</p>")
            old = self.make_material(
                old_path, uid="canvas:syllabus-file:1:old"
            )
            first = CourseMaterialsSync(
                FakeCanvas([old]),
                FakeAI([deadline()]),
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})
            self.assertTrue(first.complete)

            current_path = root / "current.html"
            current_path.write_text("<p>Current course outline</p>")
            current = self.make_material(
                current_path, uid="canvas:syllabus-page:1"
            )
            second = CourseMaterialsSync(
                FakeCanvas(
                    [current],
                    warnings=("Could not retrieve syllabus files",),
                    incomplete_course_ids=(1,),
                ),
                FakeAI([]),
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})

        self.assertTrue(second.complete)
        self.assertEqual(len(second.items), 1)
        self.assertTrue(any("cached deadlines retained" in item for item in second.warnings))

    def test_complete_scan_retires_removed_cached_source(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            old_path = root / "old.html"
            old_path.write_text("<p>Midterm Exam: October 20, 2026 at 1:30 PM</p>")
            old = self.make_material(
                old_path, uid="canvas:syllabus-file:1:old"
            )
            CourseMaterialsSync(
                FakeCanvas([old]),
                FakeAI([deadline()]),
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})

            current_path = root / "current.html"
            current_path.write_text("<p>Current course outline</p>")
            current = self.make_material(
                current_path, uid="canvas:syllabus-page:1"
            )
            report = CourseMaterialsSync(
                FakeCanvas([current]),
                FakeAI([]),
                index_path=root / "index.json",
                now_provider=lambda: NOW,
            ).sync(active_course_ids={1})

            index_text = (root / "index.json").read_text()

        self.assertTrue(report.complete)
        self.assertEqual(report.items, ())
        self.assertNotIn("canvas:syllabus-file:1:old", index_text)


if __name__ == "__main__":
    unittest.main()
