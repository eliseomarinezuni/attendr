from __future__ import annotations

import json
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from google.genai import errors

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from academic_assistant import (
    AIAssistant,
    AIConfigurationError,
    AIInputError,
    AIProviderError,
    extract_pdf_text_chunks,
    extract_powerpoint_text_chunks,
    hybrid_quiz_discord_payload,
    quiz_discord_payload,
    send_quiz_to_discord,
)
from academic_assistant.ai_efficiency import AI_USAGE


def question(number: int) -> dict:
    return {
        "question": f"Conceptual question {number}?",
        "options": [
            f"Option A{number}",
            f"Option B{number}",
            f"Option C{number}",
            f"Option D{number}",
        ],
        "correct_answer": "B",
        "explanation": f"B is correct for conceptual reason {number}.",
    }


class FakeModels:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def generate_content(self, **kwargs):
        self.calls.append(kwargs)
        response = self.responses.pop(0)
        if isinstance(response, BaseException):
            raise response
        return SimpleNamespace(text=response)


class FakeInteractions:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        response = self.responses.pop(0)
        if isinstance(response, BaseException):
            raise response
        return SimpleNamespace(output_text=response)


class FakeClient:
    def __init__(self, responses):
        self.models = FakeModels(responses)
        self.interactions = FakeInteractions(responses)


class FakeNotifier:
    def __init__(self):
        self.calls = []

    def send_custom_notification(
        self, event_key, payload, *, force=False, destination="announcements"
    ):
        self.calls.append((event_key, payload, force, destination))
        return True


class ProviderFailure(Exception):
    def __init__(self, status_code, retry_after=None):
        super().__init__(f"provider failure {status_code}")
        self.status_code = status_code
        self.response = SimpleNamespace(
            status_code=status_code,
            headers={} if retry_after is None else {"Retry-After": retry_after},
        )


class AIAssistantTests(unittest.TestCase):
    def setUp(self):
        AI_USAGE.reset()

    def test_rate_limit_retries_then_succeeds_without_real_sleep(self):
        response = json.dumps([question(1), question(2), question(3)])
        fake = FakeClient([ProviderFailure(429, "3"), response])
        sleeps = []
        assistant = AIAssistant(
            "test-key",
            client=fake,
            sleeper=sleeps.append,
            random_provider=lambda: 0.0,
        )

        result = assistant.generate_quiz("Trees", 3)

        self.assertEqual(len(result), 3)
        self.assertEqual(sleeps, [3.0])
        self.assertEqual(len(fake.models.calls), 2)

    def test_repeated_rate_limit_stops_at_bound_and_opens_run_circuit(self):
        fake = FakeClient([ProviderFailure(429)] * 4)
        sleeps = []
        assistant = AIAssistant(
            "test-key",
            client=fake,
            sleeper=sleeps.append,
            random_provider=lambda: 0.0,
        )

        with self.assertRaises(AIProviderError) as raised:
            assistant.generate_quiz("Trees", 3)
        self.assertTrue(raised.exception.transient)
        self.assertEqual(sleeps, [1.0, 2.0, 4.0])
        self.assertEqual(len(fake.models.calls), 4)

        with self.assertRaises(AIProviderError):
            assistant.generate_quiz("Another source", 3)
        self.assertEqual(len(fake.models.calls), 4)

    def test_permanent_client_error_is_not_retried(self):
        fake = FakeClient([ProviderFailure(400)])
        sleeps = []
        assistant = AIAssistant("test-key", client=fake, sleeper=sleeps.append)

        with self.assertRaises(AIProviderError) as raised:
            assistant.generate_quiz("Trees", 3)

        self.assertFalse(raised.exception.transient)
        self.assertEqual(len(fake.models.calls), 1)
        self.assertEqual(sleeps, [])

    def test_transient_server_error_retries_then_succeeds(self):
        response = json.dumps([question(1), question(2), question(3)])
        fake = FakeClient([ProviderFailure(500), response])
        sleeps = []
        assistant = AIAssistant(
            "test-key",
            client=fake,
            sleeper=sleeps.append,
            random_provider=lambda: 0.0,
        )

        self.assertEqual(len(assistant.generate_quiz("Trees", 3)), 3)
        self.assertEqual(sleeps, [1.0])
        self.assertEqual(len(fake.models.calls), 2)

    def test_request_governor_serializes_multiple_sources_without_real_sleep(self):
        class Clock:
            value = 0.0

            def now(self):
                return self.value

            def sleep(self, seconds):
                self.value += seconds

        clock = Clock()
        response = json.dumps([question(1), question(2), question(3)])
        fake = FakeClient([response, response])
        assistant = AIAssistant(
            "test-key",
            client=fake,
            sleeper=clock.sleep,
            monotonic_provider=clock.now,
            minimum_request_interval=2.0,
        )

        assistant.generate_quiz("Trees", 3)
        assistant.generate_quiz("Graphs", 3)

        self.assertEqual(clock.value, 2.0)
        self.assertEqual(len(fake.models.calls), 2)

    def test_usage_metrics_count_requests_chars_retries_and_rate_limits(self):
        response = json.dumps([question(1), question(2), question(3)])
        assistant = AIAssistant(
            "secret-that-must-not-appear",
            client=FakeClient([ProviderFailure(429), response]),
            sleeper=lambda _: None,
            random_provider=lambda: 0.0,
        )

        with self.assertLogs("attendr.ai", level="WARNING") as logs:
            assistant.generate_quiz("Trees", 3)

        usage = AI_USAGE.snapshot()["quiz_generation"]
        self.assertEqual((usage.requests, usage.retries, usage.rate_limits), (2, 1, 1))
        self.assertGreater(usage.input_chars, 0)
        self.assertNotIn("secret-that-must-not-appear", " ".join(logs.output))

    def test_task_specific_model_can_be_configured(self):
        fake = FakeClient([json.dumps([question(1), question(2), question(3)])])
        assistant = AIAssistant(
            "test-key",
            client=fake,
            task_models={"quiz_generation": "gemini-quiz-test"},
        )

        assistant.generate_quiz("Trees", 3)

        self.assertEqual(fake.models.calls[0]["model"], "gemini-quiz-test")

    def test_hybrid_quiz_contains_two_mcq_and_one_short_answer(self):
        response = [
            {**question(1), "question_type": "multiple_choice"},
            {**question(2), "question_type": "multiple_choice"},
            {
                "question_type": "short_answer",
                "question": "Explain the invariant.",
                "options": [],
                "correct_answer": "The invariant remains true after every operation.",
                "explanation": "It supports the correctness proof.",
            },
        ]
        fake = FakeClient([json.dumps(response)])
        assistant = AIAssistant("test-key", client=fake)
        result = assistant.generate_hybrid_quiz("Lecture content")
        payload = hybrid_quiz_discord_payload("Algorithms", result)
        self.assertEqual([item["question_type"] for item in result], [
            "multiple_choice", "multiple_choice", "short_answer"
        ])
        self.assertEqual(len(payload["embeds"]), 3)

    def test_powerpoint_text_extraction(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "lecture.pptx"
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = "Binary Search Trees"
            slide.placeholders[1].text = "The left subtree contains smaller keys."
            deck.save(path)
            chunks = extract_powerpoint_text_chunks(path)
        self.assertIn("Binary Search Trees", chunks[0].text)
    def test_generate_quiz_uses_structured_schema_and_returns_exact_count(self):
        fake = FakeClient([json.dumps([question(1), question(2), question(3)])])
        assistant = AIAssistant("test-key", client=fake)

        result = assistant.generate_quiz("Binary Search Trees", 3)

        self.assertEqual(len(result), 3)
        self.assertEqual(result[0]["correct_answer"], "B")
        call = fake.models.calls[0]
        self.assertEqual(call["model"], "gemini-3.6-flash")
        self.assertEqual(call["config"].response_mime_type, "application/json")
        self.assertIsNotNone(call["config"].response_schema)
        self.assertNotIn("test-key", call["contents"])

    def test_aq_key_uses_interactions_structured_output(self):
        fake = FakeClient(
            [json.dumps({"results": [question(1), question(2), question(3)]})]
        )
        assistant = AIAssistant("AQ.test-key", client=fake)

        result = assistant.generate_quiz("Binary Search Trees", 3)

        self.assertEqual(len(result), 3)
        self.assertEqual(fake.models.calls, [])
        call = fake.interactions.calls[0]
        self.assertNotIn("response_mime_type", call)
        self.assertEqual(call["response_format"]["type"], "text")
        self.assertIn("schema", call["response_format"])
        schema = call["response_format"]["schema"]
        self.assertNotIn("$defs", json.dumps(schema))
        self.assertNotIn("$ref", json.dumps(schema))

    def test_aq_key_accepts_text_from_interaction_outputs(self):
        fake = FakeClient([])
        fake.interactions.create = lambda **kwargs: SimpleNamespace(
            output_text=None,
            outputs=[
                SimpleNamespace(
                    text=json.dumps(
                        {"results": [question(1), question(2), question(3)]}
                    )
                )
            ],
        )
        assistant = AIAssistant("AQ.test-key", client=fake)

        result = assistant.generate_quiz("Binary Search Trees", 3)

        self.assertEqual(len(result), 3)

    def test_generate_quiz_rejects_wrong_count_and_invalid_options(self):
        fake = FakeClient([json.dumps([question(1), question(2)])])
        assistant = AIAssistant("test-key", client=fake)

        with self.assertRaisesRegex(AIProviderError, "expected 3"):
            assistant.generate_quiz("Photosynthesis", 3)

        with self.assertRaisesRegex(AIInputError, "3, 4, or 5"):
            assistant.generate_quiz("Photosynthesis", 2)

    def test_generated_objects_with_unexpected_keys_are_rejected(self):
        invalid = [question(1), question(2), question(3)]
        invalid[0]["unexpected"] = "must not pass validation"
        assistant = AIAssistant("test-key", client=FakeClient([json.dumps(invalid)]))

        with self.assertRaisesRegex(AIProviderError, "invalid structured data"):
            assistant.generate_quiz("Binary Search Trees", 3)

    def test_extract_syllabus_schedule_returns_clean_json_shape(self):
        response = [
            {
                "week_or_date": "Week 1",
                "topic": "Introduction",
                "readings_or_tasks": ["Chapter 1"],
                "major_deadlines": [],
            }
        ]
        fake = FakeClient([f"```json\n{json.dumps(response)}\n```"])
        assistant = AIAssistant("test-key", client=fake)

        result = assistant.extract_syllabus_schedule("Week 1: Introduction")

        self.assertEqual(result, response)

    def test_clear_deadline_bypasses_gemini_deterministically(self):
        response = [
            {
                "title": "Midterm Exam",
                "due_date": "2026-10-20",
                "due_time": "13:30",
                "kind": "exam",
                "source_evidence": "Midterm Exam: October 20 at 1:30 PM",
            }
        ]
        fake = FakeClient([json.dumps(response)])
        assistant = AIAssistant("test-key", client=fake)

        result = assistant.extract_major_deadlines(
            "Midterm Exam: October 20, 2026 at 1:30 PM",
            course_name="Algorithms",
            source_title="Course syllabus",
        )

        self.assertEqual(result[0]["due_date"], "2026-10-20")
        self.assertEqual(result[0]["due_time"], "13:30")
        self.assertEqual(fake.models.calls, [])

    def test_extract_major_deadlines_supplies_verified_timetable_context(self):
        response = [
            {
                "title": "Quiz 1",
                "due_date": "2026-10-09",
                "due_time": "15:40",
                "kind": "quiz",
                "source_evidence": "October 8 & 9, 2026; Time: Lecture Time",
            }
        ]
        fake = FakeClient([json.dumps(response)])
        assistant = AIAssistant("test-key", client=fake)

        result = assistant.extract_major_deadlines(
            "Quiz 1: October 8 & 9, 2026; Time: Lecture Time",
            course_name="Algorithms",
            source_title="Course syllabus",
            schedule_context="Lecture: Friday 15:40-17:00.",
        )

        prompt = fake.models.calls[0]["contents"]
        self.assertEqual(result, response)
        self.assertIn("VERIFIED STUDENT TIMETABLE", prompt)
        self.assertIn("Lecture: Friday 15:40-17:00", prompt)
        self.assertIn("Do not return the other section's date", prompt)

    def test_deadline_extraction_sends_only_structurally_relevant_context(self):
        response = [{
            "title": "Quiz 1",
            "due_date": "2026-10-09",
            "due_time": None,
            "kind": "quiz",
            "source_evidence": "Quiz 1: October 8 & 9, 2026",
        }]
        fake = FakeClient([json.dumps(response)])
        assistant = AIAssistant("test-key", client=fake, max_input_chars=2_000)
        source = "\n".join(
            [*(f"Policy section {index} " + "x" * 100 for index in range(200)),
             "Assessment Schedule:", "Quiz 1: October 8 & 9, 2026"]
        )

        assistant.extract_major_deadlines(
            source,
            course_name="Algorithms",
            source_title="Course syllabus",
            current_date=date(2026, 9, 1),
        )

        prompt = fake.models.calls[0]["contents"]
        self.assertIn("Quiz 1: October 8 & 9, 2026", prompt)
        self.assertNotIn("Policy section 199", prompt)
        self.assertLess(len(prompt), len(source) / 5)

    def test_extract_major_deadlines_rejects_impossible_dates(self):
        response = [
            {
                "title": "Final Exam",
                "due_date": "2026-02-30",
                "due_time": None,
                "kind": "exam",
                "source_evidence": "Final Exam: February 30",
            }
        ]
        assistant = AIAssistant("test-key", client=FakeClient([json.dumps(response)]))

        with self.assertRaisesRegex(AIProviderError, "invalid structured data"):
            assistant.extract_major_deadlines(
                "Final Exam: February 30, 2026",
                course_name="Algorithms",
                source_title="Course syllabus",
            )

    def test_missing_key_and_oversized_material_fail_before_api_call(self):
        with self.assertRaisesRegex(AIConfigurationError, "GEMINI_API_KEY"):
            AIAssistant("")

        fake = FakeClient([])
        assistant = AIAssistant("test-key", max_input_chars=1_000, client=fake)
        with self.assertRaisesRegex(AIInputError, "safe limit"):
            assistant.generate_quiz("x" * 1_001)
        self.assertEqual(fake.models.calls, [])

    def test_provider_errors_are_classified_without_raw_details(self):
        provider_error = errors.ClientError(
            400,
            {
                "error": {
                    "message": "API key not valid: secret-value",
                    "status": "INVALID_ARGUMENT",
                }
            },
        )

        safe_error = AIAssistant._safe_provider_error(provider_error, "quiz generation")

        self.assertIn("rejected GEMINI_API_KEY", str(safe_error))
        self.assertNotIn("secret-value", str(safe_error))

    def test_pdf_extraction_cleans_and_chunks_consecutive_pages(self):
        pages = [
            SimpleNamespace(extract_text=lambda: "Binary search\n\ntrees " + "a" * 600),
            SimpleNamespace(extract_text=lambda: "Balancing\n\nrotations " + "b" * 600),
        ]
        fake_reader = SimpleNamespace(is_encrypted=False, pages=pages)
        with (
            tempfile.NamedTemporaryFile(suffix=".pdf") as pdf_file,
            patch(
                "academic_assistant.ai_assistant.PdfReader", return_value=fake_reader
            ),
        ):
            chunks = extract_pdf_text_chunks(
                pdf_file.name, max_chars=1_000, overlap_chars=100
            )

        self.assertEqual(len(chunks), 2)
        self.assertEqual(chunks[0].page_start, 1)
        self.assertEqual(chunks[1].page_end, 2)
        self.assertIn("Binary search", chunks[0].text)

    def test_discord_payload_hides_answers_and_dispatches_through_notifier(self):
        questions = [question(1), question(2), question(3)]
        questions[0]["explanation"] = "Never expose ||nested spoiler|| markers."

        payload = quiz_discord_payload("Binary Search Trees", questions)

        self.assertEqual(len(payload["embeds"]), 3)
        answer = payload["embeds"][0]["fields"][1]["value"]
        self.assertTrue(answer.startswith("||"))
        self.assertTrue(answer.endswith("||"))
        self.assertIn("Answer: B", answer)
        self.assertIn("｜｜nested spoiler｜｜", answer)
        self.assertEqual(payload["allowed_mentions"], {"parse": []})

        notifier = FakeNotifier()
        sent = send_quiz_to_discord(
            notifier,
            "Binary Search Trees",
            questions,
            force=True,
            event_key="test-quiz",
        )
        self.assertTrue(sent)
        self.assertEqual(notifier.calls[0][0], "test-quiz")
        self.assertTrue(notifier.calls[0][2])
        self.assertEqual(notifier.calls[0][3], "lecture_quizzes")


if __name__ == "__main__":
    unittest.main()


def test_provider_timeout_has_safe_actionable_reason():
    timeout = type("APITimeoutError", (Exception,), {})("secret provider request")
    error = AIAssistant._safe_provider_error(timeout, "hybrid quiz generation")
    assert "timed out" in str(error)
    assert "secret" not in str(error)


def test_gemini_client_allows_three_minutes_for_lecture_generation():
    with patch("academic_assistant.ai_assistant.genai.Client") as client:
        AIAssistant("test-key")
    assert client.call_args.kwargs["http_options"].timeout == 180_000
