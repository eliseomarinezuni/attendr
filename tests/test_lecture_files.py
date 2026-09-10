from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import Mock
import logging
import socket
import zipfile

import pytest
import requests
from pptx import Presentation

from academic_assistant.lecture_files import (
    DownloadDestination,
    LectureDownloadError,
    _PinnedHTTPSAdapter,
    powerpoint_slide_text,
    stream_download,
    validate_download_destination,
)
from academic_assistant.ai_assistant import extract_powerpoint_text_chunks
from academic_assistant.canvas_client import CanvasClient
from academic_assistant.state_store import StateStore


def deck(path):
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = 'Lecture 1'
    slide.placeholders[1].text = 'Matrices and transformations'
    slide.notes_slide.notes_text_frame.text = 'Remember multiplication order'
    p.save(path)


def test_extraction_does_not_read_embedded_media(tmp_path, monkeypatch):
    path = tmp_path / 'deck.pptx'
    deck(path)
    with zipfile.ZipFile(path, 'a') as archive:
        archive.writestr('ppt/media/huge.mp4', b'ignored media')
    original = zipfile.ZipFile.read
    def read(self, name, *args, **kwargs):
        assert not (name.filename if hasattr(name, 'filename') else name).startswith('ppt/media/')
        return original(self, name, *args, **kwargs)
    monkeypatch.setattr(zipfile.ZipFile, 'read', read)
    text = '\n'.join(c.text for c in extract_powerpoint_text_chunks(path))
    assert 'Matrices and transformations' in text
    assert 'Remember multiplication order' in text


def test_xml_limit_rejects_before_read(tmp_path, monkeypatch):
    path = tmp_path / 'deck.pptx'
    deck(path)
    monkeypatch.setattr('academic_assistant.lecture_files.MAX_XML_BYTES', 10)
    with pytest.raises(ValueError, match='extraction limit'):
        powerpoint_slide_text(path)


def response(blocks, status=200, headers=None):
    r = Mock()
    r.__enter__ = Mock(return_value=r)
    r.__exit__ = Mock(return_value=False)
    r.status_code, r.headers = status, headers or {}
    r.iter_content.return_value = iter(blocks)
    return r


def dns_result(*addresses):
    return [
        (
            socket.AF_INET6 if ':' in address else socket.AF_INET,
            socket.SOCK_STREAM,
            socket.IPPROTO_TCP,
            '',
            (address, 443, 0, 0) if ':' in address else (address, 443),
        )
        for address in addresses
    ]


def public_dns(monkeypatch, mapping=None):
    mapping = mapping or {}
    monkeypatch.setattr(
        socket,
        'getaddrinfo',
        lambda hostname, *_args, **_kwargs: dns_result(
            *mapping.get(hostname, ('93.184.216.34',))
        ),
    )


def test_stream_counts_actual_bytes_and_removes_partial(tmp_path, monkeypatch):
    public_dns(monkeypatch)
    r = response([b'123', b'456'])
    get = Mock(return_value=r)
    monkeypatch.setattr(requests.Session, 'get', get)
    path = tmp_path / 'download'
    with pytest.raises(ValueError, match='byte or time'):
        stream_download('https://canvas.test/file', path, 5)
    assert not path.exists()
    assert get.call_args.kwargs['stream'] is True
    assert 'headers' not in get.call_args.kwargs


def test_stream_rejects_insecure_redirect(tmp_path, monkeypatch):
    public_dns(monkeypatch)
    monkeypatch.setattr(requests.Session, 'get', Mock(return_value=response([], 302, {'Location': 'http://unsafe.test'})))
    with pytest.raises(LectureDownloadError, match='invalid URL'):
        stream_download('https://canvas.test/file', tmp_path / 'download', 100)


def test_stream_retries_transport_failure(tmp_path, monkeypatch):
    public_dns(monkeypatch)
    r = response([b'content'])
    get = Mock(side_effect=[requests.Timeout('secret'), r])
    monkeypatch.setattr(requests.Session, 'get', get)
    monkeypatch.setattr('academic_assistant.lecture_files.time.sleep', lambda _: None)
    path = tmp_path / 'download'
    stream_download('https://canvas.test/file', path, 100)
    assert path.read_bytes() == b'content'
    assert get.call_count == 2


def test_public_https_destination_and_signed_query_are_accepted(tmp_path, monkeypatch, caplog):
    public_dns(monkeypatch)
    get = Mock(return_value=response([b'content']))
    monkeypatch.setattr(requests.Session, 'get', get)
    signed_url = 'https://cdn.canvas.test/file?verifier=signed-private-value&download=1'
    with caplog.at_level(logging.DEBUG):
        stream_download(signed_url, tmp_path / 'download', 100)
    assert get.call_args.args[0] == signed_url
    assert 'signed-private-value' not in caplog.text


def test_public_canvas_and_cdn_redirect_chain_is_revalidated(tmp_path, monkeypatch):
    resolutions = []
    mapping = {
        'canvas.test': ('93.184.216.34',),
        'cdn.canvas.test': ('2606:4700:4700::1111',),
    }
    monkeypatch.setattr(socket, 'getaddrinfo', lambda hostname, *_args, **_kwargs: (
        resolutions.append(hostname) or dns_result(*mapping[hostname])
    ))
    get = Mock(side_effect=[
        response([], 302, {'Location': 'https://cdn.canvas.test/signed?key=value'}),
        response([b'content']),
    ])
    monkeypatch.setattr(requests.Session, 'get', get)
    path = tmp_path / 'download'
    stream_download('https://canvas.test/file', path, 100)
    assert path.read_bytes() == b'content'
    assert resolutions == ['canvas.test', 'cdn.canvas.test']
    assert get.call_count == 2


@pytest.mark.parametrize(
    ('url', 'addresses'),
    [
        ('https://127.0.0.1/file', ('127.0.0.1',)),
        ('https://localhost/file', ('127.0.0.1',)),
        ('https://localhost./file', ('127.0.0.1',)),
        ('https://10.1.2.3/file', ('10.1.2.3',)),
        ('https://172.16.2.3/file', ('172.16.2.3',)),
        ('https://192.168.2.3/file', ('192.168.2.3',)),
        ('https://169.254.169.254/latest/meta-data', ('169.254.169.254',)),
        ('https://[::1]/file', ('::1',)),
        ('https://[fc00::1]/file', ('fc00::1',)),
        ('https://[fe80::1]/file', ('fe80::1',)),
        ('https://private.test/file', ('10.0.0.8',)),
        ('https://mixed.test/file', ('93.184.216.34', '192.168.1.8')),
    ],
)
def test_unsafe_destination_ranges_are_rejected_before_request(
    url, addresses, tmp_path, monkeypatch
):
    monkeypatch.setattr(socket, 'getaddrinfo', Mock(return_value=dns_result(*addresses)))
    get = Mock()
    monkeypatch.setattr(requests.Session, 'get', get)
    with pytest.raises(LectureDownloadError, match='unsafe destination'):
        stream_download(url, tmp_path / 'download', 100)
    get.assert_not_called()


@pytest.mark.parametrize('url', ['http://public.test/file', 'https://user:password@public.test/file'])
def test_invalid_url_is_rejected_without_dns_or_request(url, tmp_path, monkeypatch):
    resolve = Mock()
    get = Mock()
    monkeypatch.setattr(socket, 'getaddrinfo', resolve)
    monkeypatch.setattr(requests.Session, 'get', get)
    with pytest.raises(LectureDownloadError, match='invalid URL'):
        stream_download(url, tmp_path / 'download', 100)
    resolve.assert_not_called()
    get.assert_not_called()


def test_dns_resolution_failure_is_rejected_without_retry(tmp_path, monkeypatch):
    resolve = Mock(side_effect=socket.gaierror('private resolver detail'))
    get = Mock()
    monkeypatch.setattr(socket, 'getaddrinfo', resolve)
    monkeypatch.setattr(requests.Session, 'get', get)
    with pytest.raises(LectureDownloadError, match='resolution failed'):
        stream_download('https://unknown.test/file', tmp_path / 'download', 100)
    assert resolve.call_count == 1
    get.assert_not_called()


def test_private_metadata_redirect_is_rejected_before_second_request(tmp_path, monkeypatch):
    mapping = {
        'canvas.test': ('93.184.216.34',),
        'metadata.test': ('169.254.169.254',),
    }
    public_dns(monkeypatch, mapping)
    get = Mock(return_value=response([], 302, {'Location': 'https://metadata.test/latest'}))
    monkeypatch.setattr(requests.Session, 'get', get)
    with pytest.raises(LectureDownloadError, match='unsafe destination'):
        stream_download('https://canvas.test/file', tmp_path / 'download', 100)
    assert get.call_count == 1


def test_redirect_chain_remains_capped(tmp_path, monkeypatch):
    public_dns(monkeypatch)
    get = Mock(return_value=response([], 302, {'Location': '/next'}))
    monkeypatch.setattr(requests.Session, 'get', get)
    with pytest.raises(ValueError, match='Too many lecture download redirects'):
        stream_download('https://canvas.test/file', tmp_path / 'download', 100)
    assert get.call_count == 6


def test_validator_accepts_all_public_dns_answers(monkeypatch):
    public_dns(monkeypatch, {'cdn.test': ('93.184.216.34', '2606:4700:4700::1111')})
    destination = validate_download_destination('https://cdn.test/file')
    assert destination.addresses == ('93.184.216.34', '2606:4700:4700::1111')


def test_connection_pool_pins_ip_but_preserves_tls_hostname():
    destination = DownloadDestination('cdn.test', 443, ('93.184.216.34',))
    adapter = _PinnedHTTPSAdapter(destination)
    request = requests.Request('GET', 'https://cdn.test/file').prepare()
    pool = adapter.get_connection_with_tls_context(request, verify=True)
    connection = pool._new_conn()
    assert connection._dns_host == '93.184.216.34'
    assert connection.server_hostname == 'cdn.test'
    assert connection.assert_hostname == 'cdn.test'
    adapter.close()


def test_large_deck_cache_survives_new_runner_and_invalidates_revision(tmp_path, monkeypatch):
    source = tmp_path / 'original.pptx'
    deck(source)
    download = Mock(side_effect=lambda url, dest, limit: dest.write_bytes(source.read_bytes()))
    monkeypatch.setattr('academic_assistant.canvas_client.stream_download', download)
    file = NS(size=518508549, updated_at='2026-09-09T00:00:00Z', url='https://canvas.test/signed')
    course = NS(id=1, name='Graphics')
    def client():
        c = object.__new__(CanvasClient)
        c.base_url = 'https://canvas.test'
        c.state_store = StateStore(tmp_path / 'state.db')
        return c
    args = (file, course, '8', 'Lecture 1', tmp_path / 'materials', 768 * 1024**2, 'Week 1', 1, 1, False)
    material = client()._large_lecture_text(*args)
    assert 'Matrices' in material.local_path.read_text()
    material.local_path.unlink()
    assert client()._large_lecture_text(*args).local_path.exists()
    assert download.call_count == 1
    file.updated_at = '2026-09-10T00:00:00Z'
    client()._large_lecture_text(*args)
    assert download.call_count == 2
    assert not list((tmp_path / 'materials').glob('lecture-*'))


def test_quiz_reads_cached_text_literally(tmp_path):
    from academic_assistant.lecture_quiz import LectureQuizRunner
    path = tmp_path / 'source.txt'
    path.write_text('Compare x < y and x > z')
    assert LectureQuizRunner._material_text(NS(local_path=path)) == path.read_text()


def test_google_slides_cache_and_login_page_rejection(tmp_path, monkeypatch):
    c = object.__new__(CanvasClient)
    c.base_url = 'https://canvas.test'
    c.state_store = StateStore(tmp_path / 'state.db')
    download = Mock(side_effect=lambda url, dest, limit: dest.write_text('Protocols and HTTP'))
    monkeypatch.setattr('academic_assistant.canvas_client.stream_download', download)
    args = (NS(id=1, name='Web'), 'https://docs.google.com/presentation/d/abc/edit',
            '01a - protocols', tmp_path, 'Basics', 2, 1)
    assert 'HTTP' in c._download_external_slides(*args).local_path.read_text()
    c._download_external_slides(*args)
    assert download.call_count == 1
    c.state_store = None
    download.side_effect = lambda url, dest, limit: dest.write_text('<!DOCTYPE html><html>Sign in</html>')
    with pytest.raises(ValueError, match='unavailable'):
        c._download_external_slides(*args)
    assert c._download_external_slides(NS(id=1), 'https://unrelated.test', '', tmp_path, '', 1, 1) is None


def test_numbered_week_labels_match_correct_session(tmp_path):
    from datetime import datetime
    from academic_assistant.course_schedule import CourseSchedule
    from academic_assistant.lecture_quiz import LectureQuizRunner
    root = Path(__file__).resolve().parents[1]
    schedule = CourseSchedule.load(root / 'data/course_schedule.json')
    session = next(s for s in schedule.sessions if s.course_key == 'web-development' and s.weekday == 1 and s.activity == 'lecture')
    runner = LectureQuizRunner(None, None, None, schedule, materials_directory=tmp_path, state_path=tmp_path / 'state.db')
    common = dict(course_name='Web Development', module_name='Basics', module_position=2, updated_at=None)
    first = NS(title='01a - protocols', item_position=1, **common)
    second = NS(title='01b - HTML', item_position=2, **common)
    ended = datetime(2026, 9, 8, 14, tzinfo=schedule.timezone)
    assert runner._select_material(session, ended, (second, first)) is first
    assert runner._select_material(session, ended, (second,)) is None
    due = schedule.ended_lecture_sessions(datetime(2026, 9, 17, 22, tzinfo=schedule.timezone), retry_hours=336)
    assert any(day.date() == ended.date() and s == session for s, day in due)


def test_introduction_module_matches_only_first_course_lecture(tmp_path):
    from datetime import datetime
    from academic_assistant.course_schedule import CourseSchedule
    from academic_assistant.lecture_quiz import LectureQuizRunner
    root = Path(__file__).resolve().parents[1]
    schedule = CourseSchedule.load(root / 'data/course_schedule.json')
    sessions = sorted(
        (s for s in schedule.sessions if s.course_key == 'computer-graphics' and s.activity == 'lecture'),
        key=lambda session: (session.weekday, session.start),
    )
    runner = LectureQuizRunner(None, None, None, schedule, materials_directory=tmp_path, state_path=tmp_path / 'state.db')
    material = NS(
        course_name='Computer Graphics & Visualization',
        module_name='Introduction',
        module_position=4,
        title='Introduction_canvas.pptx',
        item_position=4,
        updated_at=None,
    )
    first_end = datetime(2026, 9, 9, 11, tzinfo=schedule.timezone)
    second_end = datetime(2026, 9, 11, 11, tzinfo=schedule.timezone)
    assert runner._select_material(sessions[0], first_end, (material,)) is material
    assert runner._select_material(sessions[1], second_end, (material,)) is None


def test_topic_modules_map_by_canvas_order_and_bundle_same_module(tmp_path):
    from datetime import datetime
    from academic_assistant.course_schedule import CourseSchedule
    from academic_assistant.lecture_quiz import LectureQuizRunner
    root = Path(__file__).resolve().parents[1]
    schedule = CourseSchedule.load(root / 'data/course_schedule.json')
    sessions = sorted(
        (s for s in schedule.sessions if s.course_key == 'computer-graphics' and s.activity == 'lecture'),
        key=lambda item: (item.weekday, item.start),
    )
    runner = LectureQuizRunner(None, None, None, schedule, materials_directory=tmp_path, state_path=tmp_path / 'state.db')
    common = dict(course_id=1, course_name='Computer Graphics & Visualization', updated_at=None)
    intro = NS(uid='intro', source_id='1', title='Welcome', module_name='Introduction', module_position=4,
               item_position=1, module_id='40', item_id='401', **common)
    modeling_a = NS(uid='model-a', source_id='2', title='Coordinate Systems', module_name='Modeling', module_position=5,
                    item_position=1, module_id='50', item_id='501', **common)
    modeling_b = NS(uid='model-b', source_id='3', title='Transformations', module_name='Modeling', module_position=5,
                    item_position=2, module_id='50', item_id='502', **common)
    rendering = NS(uid='render', source_id='4', title='Rasterization', module_name='Rendering', module_position=6,
                   item_position=1, module_id='60', item_id='601', **common)
    ended = datetime(2026, 9, 11, 11, tzinfo=schedule.timezone)
    selected = runner._select_materials(sessions[1], ended, (rendering, modeling_b, intro, modeling_a))
    assert selected == (modeling_a, modeling_b)


def test_topic_module_order_fails_closed_when_positions_are_ambiguous(tmp_path):
    from datetime import datetime
    from academic_assistant.course_schedule import CourseSchedule
    from academic_assistant.lecture_quiz import LectureQuizRunner
    root = Path(__file__).resolve().parents[1]
    schedule = CourseSchedule.load(root / 'data/course_schedule.json')
    session = next(s for s in schedule.sessions if s.course_key == 'computer-graphics' and s.weekday == 2 and s.activity == 'lecture')
    runner = LectureQuizRunner(None, None, None, schedule, materials_directory=tmp_path, state_path=tmp_path / 'state.db')
    common = dict(course_id=1, course_name='Computer Graphics & Visualization', updated_at=None,
                  item_position=1)
    one = NS(uid='one', source_id='1', title='Vectors', module_name='Vectors', module_position=4,
             module_id='40', item_id='401', **common)
    two = NS(uid='two', source_id='2', title='Matrices', module_name='Matrices', module_position=4,
             module_id='41', item_id='402', **common)
    ended = datetime(2026, 9, 9, 11, tzinfo=schedule.timezone)
    assert runner._select_materials(session, ended, (one, two)) == ()


def test_explicit_session_mapping_overrides_module_order(tmp_path):
    from datetime import datetime
    import json
    from academic_assistant.course_schedule import CourseSchedule
    from academic_assistant.lecture_quiz import LectureQuizRunner
    root = Path(__file__).resolve().parents[1]
    data = json.loads((root / 'data/course_schedule.json').read_text())
    data['courses'][0]['lecture_materials']['sessions'] = {
        '2026-09-09': {'module_id': '60'}
    }
    schedule = CourseSchedule(data)
    session = next(s for s in schedule.sessions if s.course_key == 'computer-graphics' and s.weekday == 2 and s.activity == 'lecture')
    runner = LectureQuizRunner(None, None, None, schedule, materials_directory=tmp_path, state_path=tmp_path / 'state.db')
    common = dict(course_id=1, course_name='Computer Graphics & Visualization', updated_at=None,
                  item_position=1)
    intro = NS(uid='intro', source_id='1', title='Welcome', module_name='Introduction', module_position=4,
               module_id='40', item_id='401', **common)
    rendering = NS(uid='render', source_id='4', title='Rasterization', module_name='Rendering', module_position=6,
                   module_id='60', item_id='601', **common)
    ended = datetime(2026, 9, 9, 11, tzinfo=schedule.timezone)
    assert runner._select_materials(session, ended, (intro, rendering)) == (rendering,)


def test_stable_material_mapping_survives_title_and_position_changes(tmp_path):
    from datetime import datetime
    from academic_assistant.course_schedule import CourseSchedule
    from academic_assistant.lecture_quiz import LectureQuizRunner
    root = Path(__file__).resolve().parents[1]
    schedule = CourseSchedule.load(root / 'data/course_schedule.json')
    session = next(s for s in schedule.sessions if s.course_key == 'computer-graphics' and s.weekday == 2 and s.activity == 'lecture')
    runner = LectureQuizRunner(None, None, None, schedule, materials_directory=tmp_path, state_path=tmp_path / 'state.db')
    ended = datetime(2026, 9, 9, 11, tzinfo=schedule.timezone)
    material = NS(uid='canvas:lecture-file:1:99', source_id='99', title='Renamed Topic',
                  module_name='Moved Topic', module_position=99, item_position=8,
                  module_id='60', item_id='601', course_id=1,
                  course_name='Computer Graphics & Visualization', updated_at=None)
    runner._save_material_mapping(session.session_id(ended.date()), (material,))
    assert runner._select_materials(session, ended, (material,)) == (material,)


def test_private_slides_requires_explicit_scope(tmp_path, monkeypatch):
    import json
    from academic_assistant.google_slides import authenticated_slide_text, SlidesAccessError
    path = tmp_path / 'token.json'
    path.write_text(json.dumps({'scopes': ['https://www.googleapis.com/auth/calendar']}))
    monkeypatch.setenv('GOOGLE_SLIDES_TOKEN_FILE', str(path))
    with pytest.raises(SlidesAccessError, match='AUTH_REQUIRED'):
        authenticated_slide_text('test')


def test_private_slides_falls_back_to_authorized_reader(tmp_path, monkeypatch):
    from academic_assistant.lecture_files import LectureDownloadError
    c = object.__new__(CanvasClient)
    c.state_store = None
    monkeypatch.setattr('academic_assistant.canvas_client.stream_download', Mock(side_effect=LectureDownloadError(401)))
    reader = Mock(return_value='Authenticated lecture text')
    monkeypatch.setattr('academic_assistant.canvas_client.authenticated_slide_text', reader)
    result = c._download_external_slides(NS(id=1, name='Web'),
        'https://docs.google.com/presentation/d/abc/edit', '01a', tmp_path, '', 1, 1)
    assert result.local_path.read_text() == 'Authenticated lecture text'
    reader.assert_called_once_with('abc')


def test_teaching_week_skips_reading_week_and_starts_monday():
    from datetime import date
    from academic_assistant.course_schedule import CourseSchedule
    root = Path(__file__).resolve().parents[1]
    schedule = CourseSchedule.load(root / 'data/course_schedule.json')
    assert schedule.teaching_week(date(2026, 9, 14)) == 2
    assert schedule.teaching_week(date(2026, 10, 19)) == 6
    assert schedule.teaching_week(date(2026, 11, 30)) == 12


def test_one_bad_lecture_does_not_block_other_files(tmp_path):
    c = object.__new__(CanvasClient)
    course = NS(id=1, name='Course', course_code='')
    files = [NS(id=1, display_name='Lecture1.pdf'), NS(id=2, display_name='Lecture2.pdf')]
    resource = NS(get_modules=lambda **kw: [], get_files=lambda **kw: files)
    c._get_active_course_contexts = lambda: [NS(summary=course, resource=resource)]
    good = NS(uid='canvas:lecture-file:1:2')
    def download(*args):
        if args[2] == '1':
            raise ValueError('private body')
        return good
    c._download_lecture_file = download
    report = c.download_lecture_materials(tmp_path)
    assert report.materials == (good,)
    assert len(report.warnings) == 1
    assert 'private' not in report.warnings[0]
