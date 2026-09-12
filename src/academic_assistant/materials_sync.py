"""Canvas syllabus download, Gemini deadline extraction, and calendar normalization."""

from __future__ import annotations

import os
import re
import tempfile
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from hashlib import sha256
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Literal
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from dotenv import load_dotenv
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pydantic import BaseModel, ConfigDict, Field, ValidationError

from .ai_assistant import (
    AIAssistant,
    AIInputError,
    AIProviderError,
    MajorDeadline,
    PDFExtractionError,
    extract_pdf_text_chunks,
)
from .canvas_client import AcademicItem, CanvasClient, SyllabusMaterial
from .course_schedule import CourseSchedule
from .deadline_grounding import ground_deadline
from .state_store import StateStore

UTC = timezone.utc
EXTRACTION_VERSION = 4
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class MaterialsConfigurationError(ValueError):
    """Raised when syllabus synchronization configuration is invalid."""


class MaterialsStateError(RuntimeError):
    """Raised when the material extraction cache cannot be read or written."""


class CachedMaterial(BaseModel):
    model_config = ConfigDict(extra="forbid")

    content_sha256: str = Field(pattern=r"^[a-f0-9]{64}$")
    course_id: int
    course_name: str
    title: str
    deadlines: list[MajorDeadline]
    extraction_version: int = 1
    context_hash: str = ""


class MaterialsIndex(BaseModel):
    model_config = ConfigDict(extra="forbid")

    version: Literal[1] = 1
    sources: dict[str, CachedMaterial] = Field(default_factory=dict)


@dataclass(frozen=True, slots=True)
class MaterialsSyncReport:
    materials_found: int
    materials_analyzed: int
    cached_materials_reused: int
    items: tuple[AcademicItem, ...]
    warnings: tuple[str, ...]
    complete: bool = True
    blocking_warnings: tuple[str, ...] = ()


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self.ignored_depth += 1
        elif tag in {"br", "p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self.ignored_depth:
            self.ignored_depth -= 1
        elif tag in {"p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.ignored_depth:
            self.parts.append(data)

    def text(self) -> str:
        return "\n".join(
            line for raw in "".join(self.parts).splitlines()
            if (line := " ".join(raw.split()))
        )


class CourseMaterialsSync:
    """Synchronize syllabus sources and return grounded calendar items."""

    def __init__(
        self,
        canvas: CanvasClient,
        ai: AIAssistant,
        *,
        materials_directory: str | os.PathLike[str] = "data/materials",
        index_path: str | os.PathLike[str] = "data/materials_index.json",
        app_timezone: str = "America/Toronto",
        max_file_bytes: int = 25 * 1024 * 1024,
        future_days: int = 550,
        course_schedule: CourseSchedule | None = None,
        now_provider: Any | None = None,
        state_store: StateStore | None = None,
    ) -> None:
        if max_file_bytes < 1:
            raise MaterialsConfigurationError("Material size limit must be positive.")
        if future_days < 1:
            raise MaterialsConfigurationError(
                "Material future window must be positive."
            )
        try:
            self.timezone = ZoneInfo(app_timezone)
        except ZoneInfoNotFoundError as error:
            raise MaterialsConfigurationError(
                f"APP_TIMEZONE is not a valid IANA timezone: {app_timezone}"
            ) from error
        self.state_store = state_store
        self.canvas = canvas
        self.ai = ai
        self.materials_directory = Path(materials_directory).expanduser().resolve()
        self.index_path = Path(index_path).expanduser().resolve()
        self.max_file_bytes = max_file_bytes
        self.future_days = future_days
        self.course_schedule = course_schedule
        self._now_provider = now_provider or (lambda: datetime.now(UTC))

    @classmethod
    def from_env(
        cls,
        canvas: CanvasClient,
        env_file: str | os.PathLike[str] | None = None,
    ) -> CourseMaterialsSync:
        load_dotenv(dotenv_path=env_file, override=False)
        base_directory = Path(env_file).resolve().parent if env_file else Path.cwd()
        materials_directory = cls._resolve_path(
            os.getenv("CANVAS_MATERIALS_DIR", "data/materials"), base_directory
        )
        index_path = cls._resolve_path(
            os.getenv("CANVAS_MATERIALS_INDEX", "data/materials_index.json"),
            base_directory,
        )
        max_megabytes = cls._read_positive_int("CANVAS_MATERIAL_MAX_MB", 25)
        schedule_path = cls._resolve_path(
            os.getenv("COURSE_SCHEDULE_FILE", "data/course_schedule.json"),
            base_directory,
        )
        return cls(
            canvas,
            AIAssistant.from_env(env_file),
            state_store=StateStore(cls._resolve_path(os.getenv("ATTENDR_DB", "data/attendr.db"), base_directory)),
            materials_directory=materials_directory,
            index_path=index_path,
            app_timezone=os.getenv("APP_TIMEZONE", "America/Toronto"),
            max_file_bytes=max_megabytes * 1024 * 1024,
            future_days=cls._read_positive_int("CANVAS_MATERIAL_FUTURE_DAYS", 550),
            course_schedule=(
                CourseSchedule.load(schedule_path) if schedule_path.is_file() else None
            ),
        )

    def sync(self, *, active_course_ids: set[int] | None = None) -> MaterialsSyncReport:
        index = self._load_index()
        known_file_ids: dict[int, set[str]] = {}
        for uid, cached_source in index.sources.items():
            match = re.fullmatch(r"canvas:syllabus-file:(\d+):(\d+)", uid)
            if match:
                known_file_ids.setdefault(cached_source.course_id, set()).add(match.group(2))
        downloads = self.canvas.download_syllabus_materials(
            self.materials_directory,
            max_file_bytes=self.max_file_bytes,
            known_file_ids_by_course=known_file_ids,
        )
        analyzed = 0
        reused = 0
        changed = False
        warnings = list(downloads.warnings)
        blocking_warnings: list[str] = []
        incomplete_downloads = set(downloads.incomplete_course_ids)
        blocking_courses: set[int] = set()
        extracted: list[tuple[SyllabusMaterial, MajorDeadline]] = []
        local_today = self._now_local().date()
        provider_unavailable = False

        for material in downloads.materials:
            context_hash = sha256(str((getattr(self.ai, "model", "injected"),
                self.course_schedule.context_for_course(material.course_name) if self.course_schedule else None,
                self.timezone.key, EXTRACTION_VERSION)).encode()).hexdigest()
            cached = index.sources.get(material.uid)
            source_text: str | None = None
            deadlines: list[MajorDeadline] | None = None
            if (
                cached
                and cached.content_sha256 == material.content_sha256
                and cached.extraction_version == EXTRACTION_VERSION
                and cached.context_hash == context_hash
            ):
                try:
                    source_text = self._material_text(material)
                except (AIInputError, PDFExtractionError):
                    # A current-version entry with the same content hash was grounded
                    # before it was cached, so it remains a known-good degraded input.
                    deadlines = cached.deadlines
                else:
                    deadlines = self._ground_deadlines(
                        source_text, cached.deadlines, material, local_today,
                        warnings,
                    )
                reused += 1
            else:
                # A grounding-version bump does not require another Gemini call when
                # the exact source bytes are unchanged and every cached claim passes
                # the current deterministic verifier.
                if (
                    cached
                    and cached.content_sha256 == material.content_sha256
                    and cached.extraction_version < EXTRACTION_VERSION
                    and cached.deadlines
                ):
                    try:
                        source_text = self._material_text(material)
                    except (AIInputError, PDFExtractionError):
                        source_text = None
                    if source_text is not None:
                        migration_warnings: list[str] = []
                        migrated = self._ground_deadlines(
                            source_text,
                            cached.deadlines,
                            material,
                            local_today,
                            migration_warnings,
                        )
                        if len(migrated) == len(cached.deadlines):
                            deadlines = migrated
                            index.sources[material.uid] = CachedMaterial(
                                content_sha256=material.content_sha256,
                                course_id=material.course_id,
                                course_name=material.course_name,
                                title=material.title,
                                deadlines=deadlines,
                                extraction_version=EXTRACTION_VERSION,
                                context_hash=context_hash,
                            )
                            reused += 1
                            changed = True
                        else:
                            warnings.extend(migration_warnings)

            if deadlines is None:
                try:
                    if provider_unavailable:
                        raise AIProviderError(
                            "Gemini is temporarily unavailable for this run.",
                            transient=True,
                        )
                    if source_text is None:
                        source_text = self._material_text(material)
                    raw_deadlines = self.ai.extract_major_deadlines(
                        source_text,
                        course_name=material.course_name,
                        source_title=material.title,
                        current_date=local_today,
                        schedule_context=(
                            self.course_schedule.context_for_course(
                                material.course_name
                            )
                            if self.course_schedule
                            else None
                        ),
                    )
                    validated = [
                        MajorDeadline.model_validate(item) for item in raw_deadlines
                    ]
                    deadlines = self._ground_deadlines(
                        source_text, validated, material, local_today,
                        warnings,
                    )
                except (
                    AIInputError,
                    AIProviderError,
                    PDFExtractionError,
                    ValidationError,
                ) as error:
                    if isinstance(error, AIProviderError) and error.transient:
                        provider_unavailable = True
                    blocking_courses.add(material.course_id)
                    blocking_warning = (
                        f"Could not extract deadlines from {material.course_name}: "
                        f"{material.title}."
                    )
                    warnings.append(blocking_warning)
                    blocking_warnings.append(blocking_warning)
                    continue
                index.sources[material.uid] = CachedMaterial(
                    content_sha256=material.content_sha256,
                    course_id=material.course_id,
                    course_name=material.course_name,
                    title=material.title,
                    deadlines=deadlines,
                    extraction_version=EXTRACTION_VERSION,
                    context_hash=context_hash,
                )
                analyzed += 1
                changed = True

            for deadline in deadlines:
                extracted.append((material, deadline))

        found = {material.uid for material in downloads.materials}
        active_courses = active_course_ids if active_course_ids is not None else {value.course_id for value in index.sources.values()}
        material_courses = {material.course_id for material in downloads.materials}
        for uid, previous in list(index.sources.items()):
            if uid not in found and previous.course_id in active_courses:
                scan_incomplete = previous.course_id in incomplete_downloads
                if not scan_incomplete:
                    # A complete course scan is authoritative: the source was removed or
                    # replaced. Retiring it prevents old deadlines from living forever.
                    del index.sources[uid]
                    changed = True
                    continue

                # Canvas commonly hides the Files tab while keeping module- or
                # syllabus-linked files readable. If another current source for the
                # course was found, stale cached deadlines are a safe degraded input,
                # not grounds to freeze every calendar and study-plan update.
                if previous.course_id not in material_courses:
                    blocking_courses.add(previous.course_id)
                warnings.append(
                    f"Previously indexed syllabus temporarily unavailable: "
                    f"{previous.course_name} — {previous.title}; cached deadlines retained."
                )
                fallback = SyllabusMaterial(uid=uid, source_id=uid, course_id=previous.course_id,
                    course_name=previous.course_name, title=previous.title, content_type="text/html",
                    local_path=self.materials_directory, content_sha256=previous.content_sha256,
                    updated_at=None, html_url=None)
                if previous.extraction_version == EXTRACTION_VERSION:
                    extracted.extend((fallback, deadline) for deadline in previous.deadlines)
                else:
                    blocking_courses.add(previous.course_id)
                    blocking_warning = (
                        f"Unverified cached deadlines were not retained for "
                        f"{previous.course_name} — {previous.title}."
                    )
                    warnings.append(blocking_warning)
                    blocking_warnings.append(blocking_warning)
        if changed:
            self._save_index(index)

        items = self._normalize_deadlines(
            extracted, local_today, warnings, blocking_warnings
        )
        uncovered_courses = {
            course_id
            for course_id in incomplete_downloads
            if course_id in active_courses and course_id not in material_courses
        }
        for course_id in sorted(uncovered_courses):
            warning = (
                f"No verified syllabus source or cache is available for active course "
                f"{course_id}."
            )
            warnings.append(warning)
            blocking_warnings.append(warning)
        # Completeness answers one narrow safety question: could reconciliation
        # remove or replace a valid event because trusted source data is missing
        # or contradictory? Ordinary warnings do not affect this decision.
        complete = not (
            (blocking_courses & active_courses)
            or uncovered_courses
            or blocking_warnings
        )
        return MaterialsSyncReport(
            materials_found=len(downloads.materials),
            materials_analyzed=analyzed,
            cached_materials_reused=reused,
            items=items,
            warnings=tuple(warnings),
            complete=complete,
            blocking_warnings=tuple(blocking_warnings),
        )

    @staticmethod
    def _ground_deadlines(
        source_text: str,
        deadlines: list[MajorDeadline],
        material: SyllabusMaterial,
        reference_date: date,
        warnings: list[str],
    ) -> list[MajorDeadline]:
        grounded: list[MajorDeadline] = []
        for deadline in deadlines:
            result = ground_deadline(
                source_text, deadline, reference_date=reference_date
            )
            if result.accepted:
                grounded.append(deadline)
            else:
                locator = f", {result.locator}" if result.locator else ""
                warnings.append(
                    f"Rejected ungrounded AI deadline in {material.course_name} — "
                    f"{material.title}: {deadline.title}, expected {deadline.due_date}"
                    f"{locator} ({result.reason})."
                )
        return grounded

    def _material_text(self, material: SyllabusMaterial) -> str:
        if material.content_type == "application/pdf":
            chunks = extract_pdf_text_chunks(material.local_path)
            return "\n\n".join(chunk.text for chunk in chunks)
        if material.content_type == DOCX_CONTENT_TYPE:
            return self._extract_docx_text(material)
        try:
            raw = material.local_path.read_text(encoding="utf-8")
        except OSError as error:
            raise AIInputError(
                f"Could not read downloaded material: {material.title}"
            ) from error
        parser = _HTMLTextExtractor()
        parser.feed(raw)
        parser.close()
        text = parser.text()
        if not text:
            raise AIInputError(
                f"Downloaded material contains no text: {material.title}"
            )
        return text

    @staticmethod
    def _extract_docx_text(material: SyllabusMaterial) -> str:
        try:
            document = Document(material.local_path)
        except (OSError, ValueError, KeyError, PackageNotFoundError) as error:
            raise AIInputError(
                f"Could not read downloaded Word syllabus: {material.title}"
            ) from error
        blocks: list[str] = []
        blocks.extend(
            text
            for paragraph in document.paragraphs
            if (text := " ".join(paragraph.text.split()))
        )
        for table_number, table in enumerate(document.tables, start=1):
            blocks.append(f"[Table {table_number}]")
            for row in table.rows:
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                row_text = " | ".join(cell for cell in cells if cell)
                if row_text:
                    blocks.append(row_text)
        for section in document.sections:
            for container in (section.header, section.footer):
                blocks.extend(
                    text
                    for paragraph in container.paragraphs
                    if (text := " ".join(paragraph.text.split()))
                )
        text = "\n".join(blocks).strip()
        if not text:
            raise AIInputError(
                f"Downloaded Word syllabus contains no text: {material.title}"
            )
        return text

    def _normalize_deadlines(
        self,
        extracted: list[tuple[SyllabusMaterial, MajorDeadline]],
        local_today: date,
        warnings: list[str],
        blocking_warnings: list[str] | None = None,
    ) -> tuple[AcademicItem, ...]:
        blocking_warnings = blocking_warnings if blocking_warnings is not None else []
        grouped: dict[str, list[tuple[SyllabusMaterial, MajorDeadline]]] = {}
        last_date = local_today + timedelta(days=self.future_days)
        earliest_date = local_today - timedelta(days=1)

        for material, deadline in extracted:
            deadline_date = date.fromisoformat(deadline.due_date)
            if not earliest_date <= deadline_date <= last_date:
                continue
            uid = self._deadline_uid(material.course_id, deadline.title)
            grouped.setdefault(uid, []).append((material, deadline))

        items: list[AcademicItem] = []
        for uid, candidates in grouped.items():
            unique_dates = {deadline.due_date for _, deadline in candidates}
            if len(unique_dates) > 1:
                material, deadline = candidates[0]
                warning = (
                    f"Conflicting syllabus dates for {material.course_name}: "
                    f"{deadline.title}; calendar event skipped."
                )
                warnings.append(warning)
                blocking_warnings.append(warning)
                continue
            material, deadline = max(
                candidates, key=lambda candidate: candidate[1].due_time is not None
            )
            deadline_date = date.fromisoformat(deadline.due_date)
            effective_date = deadline_date
            lecture = (
                self.course_schedule.lecture_for_course_on(
                    material.course_name, deadline_date
                )
                if self.course_schedule and deadline.kind == "exam"
                else None
            )
            if deadline.due_time:
                parsed_time = time.fromisoformat(deadline.due_time)
            elif lecture is not None:
                parsed_time = lecture.start
            else:
                parsed_time = time.min
            due_local = datetime.combine(
                effective_date, parsed_time, tzinfo=self.timezone
            )
            all_day = deadline.due_time is None and lecture is None
            duration = timedelta(hours=2 if deadline.kind == "exam" else 1)
            end_local = (
                datetime.combine(deadline_date, lecture.end, tzinfo=self.timezone)
                if lecture is not None and deadline.due_time is None
                else due_local + duration
            )
            normalized_kind = (
                deadline.kind
                if deadline.kind in {"exam", "quiz", "assignment"}
                else "assignment"
            )
            items.append(
                AcademicItem(
                    uid=uid,
                    source="syllabus_deadline",
                    source_id=material.source_id,
                    course_id=material.course_id,
                    course_name=material.course_name,
                    title=deadline.title,
                    kind=normalized_kind,
                    due_at=due_local.astimezone(UTC),
                    due_at_local=due_local,
                    end_at=None if all_day else end_local.astimezone(UTC),
                    all_day=all_day,
                    html_url=material.html_url,
                    updated_at=material.updated_at,
                    points_possible=None,
                    submission_types=(),
                    description_html=(
                        f"Extracted from {material.title}. "
                        f"Evidence: {deadline.source_evidence}"
                    ),
                )
            )
        return tuple(sorted(items, key=lambda item: (item.due_at, item.uid)))

    def _load_index(self) -> MaterialsIndex:
        if self.state_store:
            cached = self.state_store.cache_get("materials")
            if cached is not None:
                return MaterialsIndex.model_validate(cached)
        if not self.index_path.exists():
            return MaterialsIndex()
        try:
            return MaterialsIndex.model_validate_json(
                self.index_path.read_text(encoding="utf-8")
            )
        except (OSError, ValidationError) as error:
            raise MaterialsStateError(
                f"Material extraction cache is invalid: {self.index_path}"
            ) from error

    def _save_index(self, index: MaterialsIndex) -> None:
        if self.state_store:
            self.state_store.cache_set("materials", index.model_dump(mode="json"))
            return
        self.index_path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path: Path | None = None
        try:
            descriptor, raw_path = tempfile.mkstemp(
                prefix=f".{self.index_path.name}.",
                dir=self.index_path.parent,
                text=True,
            )
            temporary_path = Path(raw_path)
            with os.fdopen(descriptor, "w", encoding="utf-8") as output:
                output.write(index.model_dump_json(indent=2))
                output.write("\n")
            temporary_path.replace(self.index_path)
        except OSError as error:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
            raise MaterialsStateError(
                f"Material extraction cache could not be saved: {self.index_path}"
            ) from error

    def _now_local(self) -> datetime:
        current = self._now_provider()
        if not isinstance(current, datetime) or current.tzinfo is None:
            raise MaterialsConfigurationError(
                "now_provider must return a timezone-aware datetime."
            )
        return current.astimezone(self.timezone)

    @staticmethod
    def _deadline_uid(course_id: int, title: str) -> str:
        normalized_title = re.sub(r"[^a-z0-9]+", " ", title.casefold()).strip()
        digest = sha256(normalized_title.encode("utf-8")).hexdigest()[:24]
        return f"canvas:material-deadline:{course_id}:{digest}"

    @staticmethod
    def _resolve_path(value: str, base_directory: Path) -> Path:
        path = Path(value).expanduser()
        return path if path.is_absolute() else base_directory / path

    @staticmethod
    def _read_positive_int(name: str, default: int) -> int:
        raw = os.getenv(name)
        if raw is None or not raw.strip():
            return default
        try:
            value = int(raw)
        except ValueError as error:
            raise MaterialsConfigurationError(f"{name} must be positive.") from error
        if value < 1:
            raise MaterialsConfigurationError(f"{name} must be positive.")
        return value
